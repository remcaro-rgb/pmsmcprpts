"""Spreadsheet, presentation and report builders.

Every number placed in a document is read from a stored dataset (by dataset_id and column),
never typed in by the model; the model supplies structure and commentary only.
"""

import io
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any, Literal

from pydantic import BaseModel, Field

from .datasets import Dataset, DatasetStore
from .formats import KINDS, Kind, check_scale, format_headline, format_value, infer_kind, points_unit, to_date
from .i18n import ACCENTS_ES, COLUMNS_ES, Lang, fmt_date
from .i18n import text as i18n_text

ACRONYMS = {"adr": "ADR", "revpar": "RevPAR", "fnb": "F&B", "id": "ID", "kpi": "KPI", "pct": "%"}

# Palette (reference data-viz chrome): ink, secondary, muted, hairline, header fill, accent.
INK, SECONDARY, MUTED = "0B0B0B", "52514E", "898781"
HAIRLINE, HEADER_FILL, ACCENT = "E1E0D9", "F0EFEC", "2A78D6"


class DocError(ValueError):
    pass


# ---------- tool input models ----------

class KpiSpec(BaseModel):
    label: str = Field(description="Tile label, sentence case, e.g. 'Occupancy'.")
    dataset_id: str
    column: str = Field(description="Column holding the value.")
    row: int = Field(0, description="0-based row index in the dataset.")
    format: Literal[KINDS] | None = Field(None, description="Override the inferred format; percent and points take fractions (0.034 shows as 3.4%).")
    note: str | None = Field(None, description="Short context under the value, e.g. 'Q3 2026, all hotels'.")


class TableSpec(BaseModel):
    dataset_id: str
    columns: list[str] | None = Field(None, description="Columns to include, in order. Default: all.")
    max_rows: int = Field(12, ge=1, le=50)
    formats: dict[str, Literal[KINDS]] | None = Field(
        None, description="Per-column format overrides; percent and points take fractions (0.034 shows as 3.4%).")


class Slide(BaseModel):
    kind: Literal["kpis", "chart", "table", "bullets"]
    title: str
    kpis: list[KpiSpec] | None = Field(None, description="kind=kpis: 1-4 tiles.")
    chart_id: str | None = Field(None, description="kind=chart: chart_id from make_chart.")
    table: TableSpec | None = Field(None, description="kind=table.")
    bullets: list[str] | None = Field(
        None, description="kind=bullets: 2-6 short points. With kind=chart: optional takeaways beside the chart.")


class Section(BaseModel):
    heading: str | None = None
    text: str | None = Field(None, description="Paragraphs separated by blank lines; lines starting with '- ' are bullets.")
    kpis: list[KpiSpec] | None = None
    chart_id: str | None = None
    table: TableSpec | None = None


class SheetSpec(BaseModel):
    name: str = Field(description="Sheet tab name (max 31 characters).")
    dataset_id: str
    columns: list[str] | None = None
    formats: dict[str, Literal[KINDS]] | None = None


# ---------- shared helpers ----------

def pretty(column: str, lang: Lang = "en") -> str:
    if lang == "es" and column.lower() in COLUMNS_ES:
        return COLUMNS_ES[column.lower()]
    words = column.split("_")
    if lang == "es":
        words = [ACCENTS_ES.get(w.lower(), w) if w.islower() else w for w in words]
    out = [ACRONYMS.get(w.lower(), w) for w in words]
    if out and out[0] == words[0]:
        out[0] = out[0][:1].upper() + out[0][1:]
    return " ".join(out)


@dataclass
class Table:
    headers: list[str]
    kinds: list[Kind]
    rows: list[list[Any]]  # raw values
    total_rows: int
    source: str
    lang: Lang = "en"

    def formatted(self) -> list[list[str]]:
        return [[format_value(v, k, self.lang) for v, k in zip(r, self.kinds)] for r in self.rows]

    def showing(self) -> str:
        return i18n_text("showing", self.lang, shown=f"{len(self.rows):,}", total=f"{self.total_rows:,}")

    @property
    def numeric(self) -> list[bool]:
        return [k in ("percent", "points", "currency", "integer", "decimal") for k in self.kinds]


def get_dataset(store: DatasetStore, dataset_id: str) -> Dataset:
    try:
        return store.get(dataset_id)
    except KeyError:
        raise DocError(f"Unknown or expired dataset_id {dataset_id!r}; run the query again.") from None


def select_table(store: DatasetStore, dataset_id: str, columns: list[str] | None,
                 formats: dict[str, str] | None, max_rows: int | None, lang: Lang = "en") -> Table:
    ds = get_dataset(store, dataset_id)
    names = [c["name"] for c in ds.columns]
    columns = columns or names
    missing = [c for c in columns if c not in names]
    if missing:
        raise DocError(f"Columns {missing} are not in dataset {dataset_id}. Columns: {names}")
    idx = [names.index(c) for c in columns]
    rows = [[r[i] for i in idx] for r in (ds.rows[:max_rows] if max_rows else ds.rows)]
    formats = formats or {}
    for c, i in zip(columns, idx):
        try:
            check_scale(c, formats.get(c), [r[i] for r in ds.rows])
        except ValueError as e:
            raise DocError(str(e)) from None
    kinds = [formats.get(c) or infer_kind(c, ds.columns[i]["type"], [r[i] for r in ds.rows])
             for c, i in zip(columns, idx)]
    return Table([pretty(c, lang) for c in columns], kinds, rows, len(ds.rows), ds.source, lang)


def kpi_value(store: DatasetStore, k: KpiSpec, lang: Lang = "en") -> str:
    ds = get_dataset(store, k.dataset_id)
    names = [c["name"] for c in ds.columns]
    if k.column not in names:
        raise DocError(f"Column {k.column!r} is not in dataset {k.dataset_id}. Columns: {names}")
    if not 0 <= k.row < len(ds.rows):
        raise DocError(f"Row {k.row} is out of range; dataset {k.dataset_id} has {len(ds.rows)} rows.")
    i = names.index(k.column)
    value = ds.rows[k.row][i]
    try:
        check_scale(k.column, k.format, [value])
    except ValueError as e:
        raise DocError(str(e)) from None
    kind = k.format or infer_kind(k.column, ds.columns[i]["type"], [r[i] for r in ds.rows])
    return format_headline(value, kind, lang)


def paragraphs(text: str) -> list[tuple[str, str]]:
    """[(kind, text)] where kind is 'p' or 'li'."""
    out = []
    for block in re.split(r"\n\s*\n", text.strip()):
        lines = block.strip().splitlines()
        if all(line.lstrip().startswith(("- ", "* ")) for line in lines):
            out += [("li", line.lstrip()[2:].strip()) for line in lines]
        else:
            out.append(("p", " ".join(line.strip() for line in lines)))
    return out


def source_line(business_date: str, lang: Lang = "en") -> str:
    d = to_date(business_date)
    return i18n_text("source", lang, date=fmt_date(d, "day", lang) if d else business_date)


# ---------- spreadsheet ----------

EXCEL_FORMATS = {"percent": "0.0%", "currency": "€#,##0.00", "integer": "#,##0", "decimal": "#,##0.00",
                 "date": "yyyy-mm-dd"}


def excel_value(value: Any, kind: Kind) -> Any:
    if kind == "date" and to_date(value):
        return to_date(value)
    if kind == "points" and isinstance(value, (int, float)):
        return value * 100  # Excel has no points format; store points with a unit suffix
    return value


def excel_points_format(lang: Lang) -> str:
    unit = points_unit(lang)
    return f'+0.0" {unit}";-0.0" {unit}";0.0" {unit}"'


def build_xlsx(store: DatasetStore, title: str, sheets: list[SheetSpec], business_date: str,
               lang: Lang = "en") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if not sheets:
        raise DocError("Provide at least one sheet.")
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    about_rows = []
    header_font = Font(bold=True, color=INK)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_border = Border(bottom=Side(style="thin", color="C3C2B7"))

    for spec in sheets:
        table = select_table(store, spec.dataset_id, spec.columns, spec.formats, None, lang)
        name = re.sub(r"[\[\]:*?/\\]", "", spec.name).strip()[:31] or i18n_text("sheet_default", lang)
        base, n = name, 2
        while name.lower() in used:
            name = f"{base[:28]} {n}"
            n += 1
        used.add(name.lower())
        ws = wb.create_sheet(name)
        ws.append(table.headers)
        for cell in ws[1]:
            cell.font, cell.fill, cell.border = header_font, header_fill, header_border
            cell.alignment = Alignment(vertical="center")
        for row in table.rows:
            ws.append([excel_value(v, k) for v, k in zip(row, table.kinds)])
        for ci, (header, kind) in enumerate(zip(table.headers, table.kinds), start=1):
            letter = get_column_letter(ci)
            number_format = (excel_points_format(lang) if kind == "points" else EXCEL_FORMATS.get(kind))
            if number_format:
                for (cell,) in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
                    cell.number_format = number_format
            if table.numeric[ci - 1]:
                ws.cell(1, ci).alignment = Alignment(horizontal="right")
            sample = [format_value(r[ci - 1], kind, lang) for r in table.rows[:200]]
            ws.column_dimensions[letter].width = min(50, max([len(header)] + [len(s) for s in sample]) + 3)
        ws.freeze_panes = "A2"
        if table.rows:
            ws.auto_filter.ref = ws.dimensions
        about_rows.append((name, len(table.rows), table.source))

    about = i18n_text("about", lang)
    while about.lower() in used:  # a data sheet already has that name
        about += " "
    ws = wb.create_sheet(about)
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([i18n_text("generated", lang), datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")])
    ws.append([i18n_text("as_of", lang), business_date])
    ws.append([])
    ws.append([i18n_text("sheet", lang), i18n_text("rows", lang), i18n_text("query", lang)])
    for cell in ws[5]:
        cell.font, cell.fill = header_font, header_fill
    for r in about_rows:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["C"].width = 100
    for (cell,) in ws.iter_rows(min_row=6, min_col=3, max_col=3):
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- presentation ----------

def build_pptx(store: DatasetStore, charts: dict[str, bytes], title: str, subtitle: str | None,
               slides: list[Slide], business_date: str, lang: Lang = "en") -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    if not slides:
        raise DocError("Provide at least one slide.")
    if len(slides) > 25:
        raise DocError("At most 25 slides.")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    rgb = RGBColor.from_string

    def plain(shape):
        """Drop the theme style (shadow/effects) from an autoshape; fill and line are set explicitly."""
        style = shape._element.find(qn("p:style"))
        if style is not None:
            shape._element.remove(style)
        return shape

    def textbox(slide, x, y, w, h, text, size, color=INK, bold=False, align=None):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, rgb(color)
        if align:
            p.alignment = align
        return tf

    def bullets(slide, x, y, w, h, items, size):
        tf = textbox(slide, x, y, w, h, "", size)
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size, p.font.color.rgb = Pt(size), rgb(INK)
            p.space_after = Pt(10)

    def frame(slide, heading, number):
        textbox(slide, 0.6, 0.4, 12.1, 0.8, heading, 28, bold=True)
        textbox(slide, 0.6, 6.95, 10, 0.35, source_line(business_date, lang), 10, MUTED)
        textbox(slide, 11.7, 6.95, 1.0, 0.35, str(number), 10, MUTED, align=PP_ALIGN.RIGHT)

    # Title slide
    s = prs.slides.add_slide(blank)
    bar = plain(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.45), Inches(1.0), Inches(0.08)))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(ACCENT)
    bar.line.fill.background()
    textbox(s, 0.8, 2.65, 11.5, 1.4, title, 40, bold=True)
    if subtitle:
        textbox(s, 0.8, 3.9, 11.5, 0.8, subtitle, 20, SECONDARY)
    textbox(s, 0.8, 6.6, 11.5, 0.4, source_line(business_date, lang), 12, MUTED)

    for number, spec in enumerate(slides, start=2):
        s = prs.slides.add_slide(blank)
        frame(s, spec.title, number)
        if spec.kind == "kpis":
            if not spec.kpis or len(spec.kpis) > 4:
                raise DocError(f"Slide {spec.title!r}: kind=kpis needs 1-4 kpis.")
            n = len(spec.kpis)
            gap, w = 0.3, (12.1 - 0.3 * (n - 1)) / n
            for i, k in enumerate(spec.kpis):
                x = 0.6 + i * (w + gap)
                tile = plain(s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(w),
                                                Inches(2.6)))
                tile.adjustments[0] = 0.06
                tile.fill.solid()
                tile.fill.fore_color.rgb = rgb("FCFCFB")
                tile.line.color.rgb = rgb(HAIRLINE)
                textbox(s, x + 0.3, 2.25, w - 0.6, 0.5, k.label, 16, SECONDARY)
                textbox(s, x + 0.3, 2.85, w - 0.6, 1.0, kpi_value(store, k, lang), 40, bold=True)
                if k.note:
                    textbox(s, x + 0.3, 3.9, w - 0.6, 0.5, k.note, 12, MUTED)
        elif spec.kind == "chart":
            png = charts.get(spec.chart_id or "")
            if png is None:
                raise DocError(f"Slide {spec.title!r}: unknown chart_id {spec.chart_id!r}; create it with make_chart.")
            if spec.bullets:
                s.shapes.add_picture(io.BytesIO(png), Inches(0.6), Inches(1.5), width=Inches(8.4))
                bullets(s, 9.3, 1.7, 3.4, 4.8, spec.bullets, 16)
            else:
                s.shapes.add_picture(io.BytesIO(png), Inches(1.95), Inches(1.35), height=Inches(5.4))
        elif spec.kind == "table":
            if not spec.table:
                raise DocError(f"Slide {spec.title!r}: kind=table needs a table.")
            t = select_table(store, spec.table.dataset_id, spec.table.columns, spec.table.formats,
                             min(spec.table.max_rows, 14), lang)
            rows, cols = len(t.rows) + 1, len(t.headers)
            shape = s.shapes.add_table(rows, cols, Inches(0.6), Inches(1.5), Inches(12.1), Inches(0.4 * rows))
            table = shape.table
            values = [t.headers] + t.formatted()
            for r in range(rows):
                for c in range(cols):
                    cell = table.cell(r, c)
                    cell.text = values[r][c]
                    p = cell.text_frame.paragraphs[0]
                    p.font.size = Pt(13 if r == 0 else 12)
                    p.font.bold = r == 0
                    p.font.color.rgb = rgb(INK)
                    if t.numeric[c]:
                        p.alignment = PP_ALIGN.RIGHT
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(HEADER_FILL if r == 0 else "FFFFFF")
            if t.total_rows > len(t.rows):
                textbox(s, 0.6, 6.55, 12, 0.35, t.showing(), 11, MUTED)
        else:
            if not spec.bullets:
                raise DocError(f"Slide {spec.title!r}: kind=bullets needs bullets.")
            bullets(s, 0.6, 1.6, 12.1, 5.0, spec.bullets, 22)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------- report (docx / pdf) ----------

def build_docx(store: DatasetStore, charts: dict[str, bytes], title: str, subtitle: str | None,
               sections: list[Section], business_date: str, lang: Lang = "en") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, 0)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.runs[0].font.size = Pt(14)
        p.runs[0].font.color.rgb = RGBColor.from_string(SECONDARY)
    meta = doc.add_paragraph(source_line(business_date, lang))
    meta.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    meta.runs[0].font.size = Pt(9)

    def shade(cell, color):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), color)
        tc_pr.append(shd)

    for sec in sections:
        if sec.heading:
            doc.add_heading(sec.heading, 1)
        if sec.kpis:
            t = doc.add_table(rows=2, cols=len(sec.kpis))
            for i, k in enumerate(sec.kpis):
                t.cell(0, i).text = k.label
                t.cell(0, i).paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(SECONDARY)
                run = t.cell(1, i).paragraphs[0].add_run(kpi_value(store, k, lang))
                run.bold, run.font.size = True, Pt(20)
                if k.note:
                    note = t.cell(1, i).add_paragraph(k.note)
                    note.runs[0].font.size = Pt(9)
                    note.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
        if sec.text:
            for kind, text in paragraphs(sec.text):
                doc.add_paragraph(text, style="List Bullet" if kind == "li" else None)
        if sec.chart_id:
            png = charts.get(sec.chart_id)
            if png is None:
                raise DocError(f"Unknown chart_id {sec.chart_id!r}; create it with make_chart.")
            doc.add_picture(io.BytesIO(png), width=Inches(6.3))
        if sec.table:
            t = select_table(store, sec.table.dataset_id, sec.table.columns, sec.table.formats, sec.table.max_rows, lang)
            table = doc.add_table(rows=len(t.rows) + 1, cols=len(t.headers))
            table.style = "Table Grid"
            for r, values in enumerate([t.headers] + t.formatted()):
                for c, value in enumerate(values):
                    cell = table.cell(r, c)
                    cell.text = value
                    para = cell.paragraphs[0]
                    para.runs[0].font.size = Pt(9)
                    para.runs[0].bold = r == 0
                    if t.numeric[c]:
                        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    if r == 0:
                        shade(cell, HEADER_FILL)
            if t.total_rows > len(t.rows):
                note = doc.add_paragraph(t.showing() + ".")
                note.runs[0].font.size = Pt(9)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


PDF_CELL_PADDING = 12  # reportlab's default 6pt left and right
PDF_CELL_SLACK = 1  # a paragraph exactly as wide as its cell still wraps


def pdf_column_widths(headers: list[str], rows: list[list[str]], width: float,
                      size: float = 8.5, min_size: float = 6.5) -> tuple[list[float], float]:
    """Column widths that keep every value and every header word on one line, at the largest
    font size (down to min_size) where the table fits the page width. Spare width goes first to
    headers that would otherwise wrap, then evenly."""
    from reportlab.pdfbase.pdfmetrics import stringWidth

    def measure(s: float) -> tuple[list[float], list[float]]:
        need, want = [], []
        for c, h in enumerate(headers):
            values = max((stringWidth(r[c], "Helvetica", s) for r in rows), default=0)
            word = max((stringWidth(w, "Helvetica-Bold", s) for w in h.split()), default=0)
            need.append(max(values, word) + PDF_CELL_PADDING + PDF_CELL_SLACK)
            want.append(max(need[-1], stringWidth(h, "Helvetica-Bold", s) + PDF_CELL_PADDING + PDF_CELL_SLACK))
        return need, want

    need, want = measure(size)
    while sum(need) > width and size > min_size:
        size = max(min_size, size - 0.5)
        need, want = measure(size)
    if sum(need) >= width:  # still too wide at the smallest size: share the width by need
        return [width * n / sum(need) for n in need], size
    spare = width - sum(need)
    extra = [w - n for w, n in zip(want, need)]
    if sum(extra) >= spare:
        return [n + spare * e / sum(extra) for n, e in zip(need, extra)], size
    left = (spare - sum(extra)) / len(headers)
    return [w + left for w in want], size


def build_pdf(store: DatasetStore, charts: dict[str, bytes], title: str, subtitle: str | None,
              sections: list[Section], business_date: str, lang: Lang = "en") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate,
                                    Spacer, Table as RLTable, TableStyle)
    from xml.sax.saxutils import escape

    ss = getSampleStyleSheet()
    ink = colors.HexColor("#" + INK)
    styles = {
        "title": ParagraphStyle("t", parent=ss["Title"], alignment=0, fontSize=24, leading=29, textColor=ink),
        "subtitle": ParagraphStyle("st", parent=ss["Normal"], fontSize=13, leading=17,
                                   textColor=colors.HexColor("#" + SECONDARY)),
        "meta": ParagraphStyle("m", parent=ss["Normal"], fontSize=8.5, textColor=colors.HexColor("#" + MUTED)),
        "h1": ParagraphStyle("h", parent=ss["Heading2"], fontSize=15, leading=19, spaceBefore=14, textColor=ink),
        "body": ParagraphStyle("b", parent=ss["Normal"], fontSize=10.5, leading=15, spaceAfter=6, textColor=ink),
        "cell": ParagraphStyle("c", parent=ss["Normal"], fontSize=8.5, leading=11),
        "cell_r": ParagraphStyle("cr", parent=ss["Normal"], fontSize=8.5, leading=11, alignment=TA_RIGHT),
        "kpi_label": ParagraphStyle("kl", parent=ss["Normal"], fontSize=9, textColor=colors.HexColor("#" + SECONDARY)),
        "kpi_value": ParagraphStyle("kv", parent=ss["Normal"], fontSize=18, leading=22, fontName="Helvetica-Bold"),
    }
    width = A4[0] - 4 * cm
    story = [Paragraph(escape(title), styles["title"])]
    if subtitle:
        story.append(Paragraph(escape(subtitle), styles["subtitle"]))
    story += [Spacer(1, 4), Paragraph(escape(source_line(business_date, lang)), styles["meta"]), Spacer(1, 10)]

    hair = colors.HexColor("#" + HAIRLINE)
    for sec in sections:
        section_start = len(story)
        if sec.heading:
            story.append(Paragraph(escape(sec.heading), styles["h1"]))
        if sec.kpis:
            cells = [[Paragraph(escape(k.label), styles["kpi_label"]) for k in sec.kpis],
                     [Paragraph(escape(kpi_value(store, k, lang)), styles["kpi_value"]) for k in sec.kpis]]
            if any(k.note for k in sec.kpis):
                cells.append([Paragraph(escape(k.note or ""), styles["meta"]) for k in sec.kpis])
            t = RLTable(cells, colWidths=[width / len(sec.kpis)] * len(sec.kpis))
            t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, hair), ("LINEAFTER", (0, 0), (-2, -1), 0.5, hair),
                                   ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
            story += [t, Spacer(1, 8)]
        if sec.text:
            items = []
            for kind, text in paragraphs(sec.text):
                if kind == "li":
                    items.append(ListItem(Paragraph(escape(text), styles["body"]), leftIndent=12))
                    continue
                if items:
                    story.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=12))
                    items = []
                story.append(Paragraph(escape(text), styles["body"]))
            if items:
                story.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=12))
        if sec.chart_id:
            png = charts.get(sec.chart_id)
            if png is None:
                raise DocError(f"Unknown chart_id {sec.chart_id!r}; create it with make_chart.")
            story += [Spacer(1, 6), Image(io.BytesIO(png), width=width, height=width * 450 / 800), Spacer(1, 6)]
        if sec.table:
            t = select_table(store, sec.table.dataset_id, sec.table.columns, sec.table.formats, sec.table.max_rows, lang)
            rows = t.formatted()
            widths, size = pdf_column_widths(t.headers, rows, width)
            cell = ParagraphStyle("c", parent=styles["cell"], fontSize=size, leading=size * 1.3)
            cell_r = ParagraphStyle("cr", parent=cell, alignment=TA_RIGHT)
            data = [[Paragraph(f"<b>{escape(h)}</b>", cell_r if n else cell) for h, n in zip(t.headers, t.numeric)]]
            data += [[Paragraph(escape(v), cell_r if n else cell) for v, n in zip(row, t.numeric)] for row in rows]
            table = RLTable(data, repeatRows=1, colWidths=widths)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + HEADER_FILL)),
                ("LINEBELOW", (0, 0), (-1, -1), 0.5, hair),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story += [table, Spacer(1, 4)]
            if t.total_rows > len(t.rows):
                story.append(Paragraph(t.showing() + ".", styles["meta"]))
        if sec.heading and len(story) - section_start > 1:
            # Keep the heading on the same page as what follows it (up to and including a chart).
            parts = story[section_start:]
            end = next((i + 1 for i, f in enumerate(parts) if isinstance(f, Image)), min(len(parts), 3))
            story[section_start:] = [KeepTogether(parts[:end])] + parts[end:]

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                      bottomMargin=2 * cm, title=title).build(story)
    return buf.getvalue()
