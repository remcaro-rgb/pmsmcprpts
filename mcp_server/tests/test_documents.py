"""Document tools, called through an in-process MCP client against the seeded database."""

import io
import os

import pytest
from mcp.client import Client

from pms_mcp.artifacts import ArtifactStore, safe_filename
from pms_mcp.documents import excel_currency_format, pretty
from pms_mcp.formats import date_labels, format_headline, format_value, infer_kind, nice_ticks
from pms_mcp.server import build_server

DSN = os.environ.get("PMS_TEST_DATABASE_URL")
needs_db = pytest.mark.skipif(not DSN, reason="PMS_TEST_DATABASE_URL not set")

MONTHLY = """SELECT hotel_code, date_trunc('month', stay_date)::date AS month,
                    round(sum(rooms_sold)::numeric / sum(rooms_available), 4) AS occupancy,
                    sum(room_revenue) AS room_revenue
             FROM v_daily_kpis WHERE stay_date BETWEEN '2026-01-01' AND '2026-06-30'
             GROUP BY 1, 2 ORDER BY 2, 1"""


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.fixture
async def ctx(tmp_path):
    async with Client(build_server(DSN, artifacts_dir=str(tmp_path))) as c:
        yield c, ArtifactStore(tmp_path)


async def call(client, tool, **args):
    r = await client.call_tool(tool, args)
    assert not r.is_error, r.content[0].text
    return r


async def error(client, tool, **args) -> str:
    r = await client.call_tool(tool, args)
    assert r.is_error
    return r.content[0].text


async def monthly(client) -> str:
    return (await call(client, "run_sql", query=MONTHLY)).structured_content["dataset_id"]


# ---------- pure helpers ----------

def test_formats():
    assert infer_kind("occupancy", "numeric", [0.81]) == "percent"
    assert infer_kind("cancellation_rate", "numeric", [0.2]) == "percent"
    assert infer_kind("rate", "numeric", [120.5]) == "currency"
    assert infer_kind("room_revenue", "numeric", [1e6]) == "currency"
    assert infer_kind("rooms_sold", "int8", [10]) == "integer"
    assert infer_kind("month", "date", ["2026-01-01"]) == "date"
    assert format_value(0.9211, "percent") == "92.1%"
    assert format_value(4083708.89, "currency") == "€4,083,709"
    assert format_value(275.314, "currency") == "€275.31"
    assert infer_kind("fully_actual", "bool", [True, False]) == "text"
    assert format_value(True, infer_kind("fully_actual", "", [True])) == "Yes"
    assert format_value(False, "text", "es") == "No"
    assert format_value(True, "text", "es") == "Sí"
    assert infer_kind("occupancy_change_pp", "numeric", [-0.0154, 0.007]) == "points"
    assert format_value(-0.0154, "points") == "-1.5 pts"
    assert format_value(0.007, "points", "es") == "+0.7 pp"
    # Spanish column names (Spanish answers name query columns in Spanish).
    assert infer_kind("ocupacion", "numeric", [0.904]) == "percent"
    assert infer_kind("tasa_cancelacion", "numeric", [0.2186]) == "percent"
    assert infer_kind("ingresos_totales", "numeric", [11127338]) == "currency"
    assert infer_kind("var_ocupacion_pp", "numeric", [-0.0047]) == "points"
    assert pretty("ocupacion", "es") == "Ocupación" and pretty("var_ocupacion_pp", "es") == "Var ocupación pp"
    # Changes already in percent (3.45 for +3.45%) are not fractions: shown as numbers, not money.
    assert infer_kind("var_adr_pct", "numeric", [3.45, 3.25]) == "decimal"
    assert infer_kind("occupancy_change_pp", "numeric", [-1.54, 0.71]) == "decimal"


def test_percent_formats_reject_values_already_in_percent():
    from pms_mcp.formats import check_scale

    check_scale("adr_change_pct", "percent", [0.0345, 0.0325])  # fractions: fine
    check_scale("occ_change_pp", "points", [-0.0154, 0.0071])
    check_scale("growth_pct", "percent", [2.5, 0.3])  # one value above 150% is plausible
    with pytest.raises(ValueError, match="already in percent"):
        check_scale("adr_change_pct", "percent", [3.45, 3.25, 3.76])
    with pytest.raises(ValueError, match="already in percent"):
        check_scale("occ_change_pp", "points", [-1.54, -0.65, 0.71])
    assert nice_ticks(0.55, 0.93) == [0.5, 0.6, 0.7, 0.8, 0.9, 1.0]
    assert date_labels(["2026-07-01", "2026-10-01"])[0] == ["Q3 2026", "Q4 2026"]
    assert date_labels(["2025-12-01", "2026-01-01"])[0] == ["Dec '25", "Jan '26"]
    assert safe_filename("Q3 report: Lisbon/Porto?", "pdf") == "Q3 report LisbonPorto.pdf"
    assert safe_filename("../../etc/passwd", "xlsx") == "etcpasswd.xlsx"


def test_currency_formats():
    assert format_value(275.314, "currency", currency="USD") == "$275.31"
    assert format_value(4083708.89, "currency", currency="gbp") == "£4,083,709"
    assert format_value(1234000, "currency", currency="COP") == "COP 1,234,000"
    assert format_value(-850.5, "currency", "es", "COP") == "COP -850.50"
    assert format_headline(11127338, "currency", currency="COP") == "COP 11.1M"
    assert format_headline(11127338, "currency") == "€11.1M"
    assert excel_currency_format("EUR") == "€#,##0.00"
    assert excel_currency_format("USD") == "$#,##0.00"
    assert excel_currency_format("COP") == '"COP "#,##0.00'


def test_artifact_ids_are_validated(tmp_path):
    store = ArtifactStore(tmp_path)
    info = store.save(b"x", "a", "png")
    assert store.read(info.file_id) == b"x"
    for bad in ("../" + info.file_id, info.file_id + ".json", "abc"):
        with pytest.raises(KeyError):
            store.read(bad)


# ---------- tools ----------

@needs_db
@pytest.mark.anyio
async def test_make_chart(ctx):
    client, files = ctx
    ds = await monthly(client)
    r = await call(client, "make_chart", dataset_id=ds, x="month", y="occupancy", series="hotel_code",
                   title="Occupancy by hotel", subtitle="Monthly, H1 2026")
    full = r.structured_content
    assert full["series"] == ["BCN", "LIS", "OPO"] and full["rows"] == 18 and full["type"] == "line"
    chart = full["chart"]
    # A Vega-Lite spec with the dataset's rows, both themes, and locales for the browser.
    spec = chart["spec"]
    assert chart["kind"] == "vega-lite" and spec["mark"]["type"] == "line" and len(spec["data"]["values"]) == 18
    assert spec["encoding"]["x"]["type"] == "temporal" and spec["encoding"]["color"]["field"] == "hotel_code"
    assert spec["data"]["values"][0]["month"].endswith("T00:00:00")  # the same day in every time zone
    assert chart["config"]["light"]["range"]["category"][0] == "#2a78d6"
    assert chart["config"]["dark"]["range"]["category"][0] == "#3987e5"
    assert chart["locale"]["time"]["months"][0] == "January"
    table = chart["table"]
    assert table["columns"] == ["Month", "BCN", "LIS", "OPO"] and table["rows"][0][0] == "January 2026"
    assert table["rows"][0][1].endswith("%")
    # The same chart as bars and areas, to switch to on the website (specs without their data).
    assert chart["type"] == "line" and set(chart["variants"]) == {"bar", "area"}
    bars = chart["variants"]["bar"]
    assert bars["mark"]["type"] == "bar" and bars["encoding"]["x"]["type"] == "ordinal" and "data" not in bars
    # The model gets the summary and the image, not the payload.
    assert '"spec"' not in r.content[0].text and full["chart_id"] in r.content[0].text
    assert r.content[1].type == "image"
    assert files.read(full["chart_id"])[:8] == b"\x89PNG\r\n\x1a\n"

    # Bars with a trend line and a second measure on its own axis.
    combo = (await call(client, "make_chart", dataset_id=ds, x="month", y="room_revenue", chart_type="bar",
                        series="hotel_code", trend="linear", y2="occupancy", title="Revenue and occupancy",
                        currency="COP")).structured_content["chart"]
    assert [list(layer) for layer in combo["spec"]["layer"]][1] == ["mark", "encoding"]
    assert combo["spec"]["resolve"] == {"scale": {"y": "independent"}}
    assert any(t.get("as") == "_trend" for layer in combo["spec"]["layer"][0]["layer"]
               for t in layer.get("transform", []))
    assert combo["locale"]["number"]["currency"] == ["COP ", ""]
    assert combo["table"]["rows"][0][1].startswith("COP ")
    assert set(combo["variants"]) == {"line", "area"}  # with the same trend and second axis
    assert combo["variants"]["line"]["resolve"] == {"scale": {"y": "independent"}}

    for kind in ("stacked_bar", "horizontal_bar", "area", "stacked_area", "boxplot", "heatmap"):
        await call(client, "make_chart", dataset_id=ds, x="month", y="occupancy", series="hotel_code",
                   chart_type=kind, title=kind)
    await call(client, "make_chart", dataset_id=ds, x="hotel_code", y="room_revenue", chart_type="donut",
               title="Share")
    await call(client, "make_chart", dataset_id=ds, x="occupancy", y="room_revenue", chart_type="scatter",
               trend="linear", title="Scatter")
    await call(client, "make_chart", dataset_id=ds, x="occupancy", chart_type="histogram", title="Spread")

    # Anything else as a Vega-Lite spec over the dataset.
    custom = await call(client, "make_chart", dataset_id=ds, title="With the average", spec={"layer": [
        {"mark": "bar", "encoding": {"x": {"field": "hotel_code", "type": "nominal"},
                                     "y": {"aggregate": "sum", "field": "room_revenue", "type": "quantitative"}}},
        {"mark": "rule", "encoding": {"y": {"aggregate": "mean", "field": "room_revenue",
                                            "type": "quantitative"}}}]})
    assert custom.structured_content["type"] == "custom"
    assert custom.structured_content["chart"]["spec"]["title"]["text"] == "With the average"


@needs_db
@pytest.mark.anyio
async def test_make_chart_rejects_bad_input(ctx):
    client, _ = ctx
    ds = await monthly(client)
    assert "not in dataset" in await error(client, "make_chart", dataset_id=ds, x="nope", y="occupancy", title="t")
    assert "not numeric" in await error(client, "make_chart", dataset_id=ds, x="month", y="hotel_code", title="t")
    assert "Unknown or expired" in await error(client, "make_chart", dataset_id="ds_x", x="a", y="b", title="t")
    nulls = (await call(client, "run_sql", query="SELECT 'a' AS x, NULL::numeric AS v")).structured_content
    assert "no values" in await error(client, "make_chart", dataset_id=nulls["dataset_id"], x="x", y="v", title="t")
    many = (await call(client, "run_sql", query="SELECT id::text AS g, 1 AS d, 2 AS v FROM guests LIMIT 13")
            ).structured_content["dataset_id"]
    assert "too many" in await error(client, "make_chart", dataset_id=many, x="d", y="v", series="g", title="t")
    # Specs carry no data of their own and load nothing from the network; fields must exist.
    assert "may not contain 'data'" in await error(client, "make_chart", dataset_id=ds, title="t",
                                                   spec={"data": {"url": "https://example.com/x.csv"},
                                                         "mark": "bar"})
    assert "may not contain 'layer[0].encoding.href'" in await error(
        client, "make_chart", dataset_id=ds, title="t",
        spec={"layer": [{"mark": "bar", "encoding": {"href": {"field": "hotel_code"}}}]})
    assert "is not a column" in await error(client, "make_chart", dataset_id=ds, title="t",
                                            spec={"mark": "bar", "encoding": {"x": {"field": "nope"}}})
    assert "could not be drawn" in await error(client, "make_chart", dataset_id=ds, title="t",
                                               spec={"mark": "no-such-mark"})
    assert "Trend lines go on" in await error(client, "make_chart", dataset_id=ds, x="hotel_code", y="occupancy",
                                              chart_type="pie", trend="linear", title="t")


@needs_db
@pytest.mark.anyio
async def test_make_spreadsheet(ctx):
    from openpyxl import load_workbook

    client, files = ctx
    ds = await monthly(client)
    r = await call(client, "make_spreadsheet", title="Occupancy H1 2026",
                   sheets=[{"name": "Monthly", "dataset_id": ds}, {"name": "Monthly", "dataset_id": ds,
                                                                     "columns": ["month", "occupancy"]}])
    f = r.structured_content["file"]
    assert f["filename"] == "Occupancy H1 2026.xlsx"
    wb = load_workbook(io.BytesIO(files.read(f["file_id"])))
    assert wb.sheetnames == ["Monthly", "Monthly 2", "About"]
    ws = wb["Monthly"]
    assert [c.value for c in ws[1]] == ["Hotel code", "Month", "Occupancy", "Room revenue"]
    assert ws.max_row == 19 and ws.freeze_panes == "A2"
    assert ws["C2"].number_format == "0.0%" and ws["D2"].number_format.startswith("€")
    assert ws["B2"].value.year == 2026
    assert "v_daily_kpis" in wb["About"]["C6"].value


@needs_db
@pytest.mark.anyio
async def test_make_presentation(ctx):
    from pptx import Presentation

    client, files = ctx
    ds = await monthly(client)
    chart = (await call(client, "make_chart", dataset_id=ds, x="month", y="occupancy", series="hotel_code",
                        title="Occupancy")).structured_content["chart_id"]
    kpis = (await call(client, "get_kpis", start="2026-01-01", end="2026-06-30", by_hotel=False,
                       granularity="year")).structured_content["dataset_id"]
    r = await call(client, "make_presentation", title="H1 2026 review", subtitle="All hotels", slides=[
        {"kind": "kpis", "title": "Headline KPIs", "kpis": [
            {"label": "Occupancy", "dataset_id": kpis, "column": "occupancy"},
            {"label": "ADR", "dataset_id": kpis, "column": "adr", "note": "H1 2026"}]},
        {"kind": "chart", "title": "Lisbon led occupancy", "chart_id": chart, "bullets": ["Spring ramp-up"]},
        {"kind": "table", "title": "Monthly detail", "table": {"dataset_id": ds, "max_rows": 5}},
        {"kind": "bullets", "title": "Next steps", "bullets": ["Push direct bookings", "Review OTA mix"]},
    ])
    prs = Presentation(io.BytesIO(files.read(r.structured_content["file"]["file_id"])))
    assert len(prs.slides) == 5
    text = lambda s: " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)  # noqa: E731
    assert "H1 2026 review" in text(prs.slides[0])
    assert "%" in text(prs.slides[1]) and "€" in text(prs.slides[1])
    assert any(sh.shape_type == 13 for sh in prs.slides[2].shapes)  # picture
    table = next(sh for sh in prs.slides[3].shapes if sh.has_table).table
    assert table.cell(0, 0).text == "Hotel code" and len(table.rows) == 6
    assert "Source: hotel PMS database" in text(prs.slides[4])

    assert "unknown chart_id" in await error(client, "make_presentation", title="t",
                                             slides=[{"kind": "chart", "title": "x", "chart_id": "nope"}])
    assert "out of range" in await error(client, "make_presentation", title="t", slides=[
        {"kind": "kpis", "title": "x", "kpis": [{"label": "a", "dataset_id": kpis, "column": "adr", "row": 9}]}])


@needs_db
@pytest.mark.anyio
@pytest.mark.parametrize("fmt", ["pdf", "docx"])
async def test_make_report(ctx, fmt):
    client, files = ctx
    ds = await monthly(client)
    chart = (await call(client, "make_chart", dataset_id=ds, x="month", y="occupancy", series="hotel_code",
                        title="Occupancy")).structured_content["chart_id"]
    r = await call(client, "make_report", title="H1 2026 performance", format=fmt, sections=[
        {"heading": "Summary", "text": "Occupancy rose through spring.\n\n- Lisbon led\n- Porto lagged",
         "kpis": [{"label": "Rows", "dataset_id": ds, "column": "occupancy"}]},
        {"heading": "Occupancy", "chart_id": chart, "table": {"dataset_id": ds, "max_rows": 6}},
    ])
    f = r.structured_content["file"]
    data = files.read(f["file_id"])
    assert f["filename"] == f"H1 2026 performance.{fmt}"
    if fmt == "pdf":
        assert data.startswith(b"%PDF") and data.count(b"/Type /Page\n") + data.count(b"/Type /Page ") >= 1
    else:
        from docx import Document
        doc = Document(io.BytesIO(data))
        texts = [p.text for p in doc.paragraphs]
        assert "Summary" in texts and "Lisbon led" in texts
        assert len(doc.inline_shapes) == 1 and len(doc.tables) == 2


@needs_db
@pytest.mark.anyio
async def test_documents_in_another_currency(ctx):
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation

    client, files = ctx
    ds = await monthly(client)
    r = await call(client, "make_spreadsheet", title="Pesos", currency="COP",
                   sheets=[{"name": "Monthly", "dataset_id": ds}])
    ws = load_workbook(io.BytesIO(files.read(r.structured_content["file"]["file_id"])))["Monthly"]
    assert ws["D2"].number_format == '"COP "#,##0.00' and ws["C2"].number_format == "0.0%"

    kpi = {"label": "Room revenue", "dataset_id": ds, "column": "room_revenue"}
    r = await call(client, "make_presentation", title="Pesos", currency="COP", slides=[
        {"kind": "kpis", "title": "Revenue", "kpis": [kpi]},
        {"kind": "table", "title": "Detail", "table": {"dataset_id": ds, "max_rows": 3}}])
    prs = Presentation(io.BytesIO(files.read(r.structured_content["file"]["file_id"])))
    tile = " ".join(sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame)
    assert "COP " in tile and "€" not in tile
    table = next(sh for sh in prs.slides[2].shapes if sh.has_table).table
    assert table.cell(1, 3).text.startswith("COP ")

    r = await call(client, "make_report", title="Pesos", format="docx", currency="usd", sections=[
        {"kpis": [kpi], "table": {"dataset_id": ds, "max_rows": 3}}])
    doc = Document(io.BytesIO(files.read(r.structured_content["file"]["file_id"])))
    kpis, detail = doc.tables
    assert kpis.cell(1, 0).text.startswith("$") and detail.cell(1, 3).text.startswith("$")
    await call(client, "make_report", title="Pesos", currency="COP", sections=[{"kpis": [kpi]}])

    assert "3-letter ISO code" in await error(client, "make_spreadsheet", title="t", currency='"; DROP',
                                              sheets=[{"name": "s", "dataset_id": ds}])


# ---------- Latin American Spanish ----------

def test_spanish_labels():
    from pms_mcp.documents import pretty, source_line

    assert date_labels(["2026-07-01", "2026-10-01"], "es")[0] == ["T3 2026", "T4 2026"]
    assert date_labels(["2025-12-01", "2026-01-01"], "es") == (["dic 25", "ene 26"],
                                                               ["diciembre de 2025", "enero de 2026"])
    assert date_labels(["2026-10-16", "2026-10-17"], "es") == (["16 oct", "17 oct"],
                                                               ["vie 16 oct 2026", "sáb 17 oct 2026"])
    assert format_value("2026-09-02", "date", "es") == "2 sept 2026"
    assert pretty("room_revenue", "es") == "Ingresos por habitaciones"
    assert pretty("hotel_code", "es") == "Hotel" and pretty("adr", "es") == "ADR"
    assert pretty("some_new_column", "es") == "Some new column"  # not in the glossary
    assert source_line("2026-10-17", "es") == "Fuente: base de datos del PMS · datos al 17 oct 2026"


@needs_db
@pytest.mark.anyio
async def test_documents_in_spanish(ctx):
    from docx import Document
    from openpyxl import load_workbook

    client, files = ctx
    ds = await monthly(client)
    chart = await call(client, "make_chart", dataset_id=ds, x="month", y="occupancy", series="hotel_code",
                       title="Ocupación por hotel", language="es")
    payload = chart.structured_content["chart"]
    assert [r[0] for r in payload["table"]["rows"][:2]] == ["enero de 2026", "febrero de 2026"]
    assert payload["table"]["columns"][0] == "Mes" and payload["locale"]["time"]["months"][0] == "enero"
    assert chart.structured_content["file"]["filename"] == "Ocupacion por hotel.png"

    r = await call(client, "make_spreadsheet", title="Ocupación S1 2026", language="es",
                   sheets=[{"name": "Mensual", "dataset_id": ds}])
    wb = load_workbook(io.BytesIO(files.read(r.structured_content["file"]["file_id"])))
    assert wb.sheetnames == ["Mensual", "Acerca de"]
    assert [c.value for c in wb["Mensual"][1]] == ["Hotel", "Mes", "Ocupación", "Ingresos por habitaciones"]
    assert wb["Acerca de"]["A2"].value == "Generado"

    r = await call(client, "make_report", title="Desempeño S1 2026", format="docx", language="es", sections=[
        {"heading": "Ocupación", "chart_id": chart.structured_content["chart_id"],
         "table": {"dataset_id": ds, "max_rows": 6}}])
    doc = Document(io.BytesIO(files.read(r.structured_content["file"]["file_id"])))
    texts = [p.text for p in doc.paragraphs]
    assert any(t.startswith("Fuente: base de datos del PMS") for t in texts)
    assert "Mostrando 6 de 18 filas." in texts
    assert doc.tables[0].cell(0, 2).text == "Ocupación"


def test_pdf_columns_fit_values_and_header_words():
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph

    from pms_mcp.documents import PDF_CELL_PADDING, pdf_column_widths

    def lines(text, width, size, bold=False):
        style = ParagraphStyle("c", fontName="Helvetica-Bold" if bold else "Helvetica", fontSize=size)
        p = Paragraph(text, style)
        p.wrap(width - PDF_CELL_PADDING, 1000)
        return len(p.blPara.lines)

    width = 482
    cases = [
        # 10 columns: the font shrinks.
        (["Hotel", "Habitaciones vendidas", "Ingresos por habitaciones", "Ingresos totales",
          "Var. ocupación pp", "Var. RevPAR %", "Var. ingresos totales %", "ADR", "RevPAR", "Ocupación"],
         ["BCN", "14,833", "€4,083,709", "€11,127,338", "-1.5 pp", "1.7%", "1.0%", "€275.31", "€246.60", "89.6%"]),
        # 9 columns at full size, where wrapping headers take the spare width (ADR wrapped here).
        (["Hotel", "Occupancy", "ADR", "RevPAR", "Room revenue", "F&B revenue", "Spa revenue", "Other revenue",
          "Total revenue"],
         ["BCN", "89.6%", "€275.31", "€246.60", "€4,083,709", "€428,281", "€244,942", "€52,981", "€4,809,913"]),
    ]
    for headers, row in cases:
        widths, size = pdf_column_widths(headers, [row], width)
        assert abs(sum(widths) - width) < 0.01 and size >= 6.5
        for c, w in enumerate(widths):  # no value and no header word is split across lines
            assert lines(row[c], w, size) == 1, (row[c], w)
            assert all(lines(word, w, size, bold=True) == 1 for word in headers[c].split()), (headers[c], w)
